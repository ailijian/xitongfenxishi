"""Read-only extraction for source audit; not needed to run the H5 app."""
from pathlib import Path
import json
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path(__file__).resolve().parents[1]
out = root / 'tmp' / 'source-audit'
out.mkdir(parents=True, exist_ok=True)
book = load_workbook(next((root / 'materials').glob('*.xlsx')), data_only=False)
for sheet in book:
    rows = [{'row': i, 'cells': [str(v) if v is not None else '' for v in row]} for i, row in enumerate(sheet.values, 1)]
    (out / (sheet.title + '.json')).write_text(json.dumps(rows, ensure_ascii=False, indent=2))
    (out / (sheet.title + '.txt')).write_text('\n'.join(f"{r['row']}: " + ' | '.join(r['cells']) for r in rows))
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(next((root / 'materials').glob('*.docx'))) as z:
    doc = ET.fromstring(z.read('word/document.xml'))
paragraphs = [''.join(p.findall('.//w:t', ns)[i].text or '' for i in range(len(p.findall('.//w:t', ns)))) for p in doc.findall('.//w:body//w:p', ns)]
(out / 'word.txt').write_text('\n'.join(f'{i}: {p}' for i,p in enumerate(paragraphs, 1)))
pdf = PdfReader(next((root / 'materials').glob('*.pdf')))
(out / 'pdf.txt').write_text('\n'.join(f'\nPAGE {i}\n' + (p.extract_text() or '') for i,p in enumerate(pdf.pages, 1)))
print(json.dumps({'sheets': [{ 'name': s.title, 'rows': s.max_row, 'columns': s.max_column } for s in book], 'wordParagraphs': len(paragraphs), 'pdfPages': len(pdf.pages)}, ensure_ascii=False))
